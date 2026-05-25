"""Dispatch a file path to the right text extractor based on extension."""
from __future__ import annotations
from pathlib import Path


class UnsupportedFileType(Exception):
    """Raised when no extractor handles this extension."""


class NotYetImplemented(Exception):
    """Raised when an extractor exists in spec but is deferred (e.g. MP4 transcription)."""


def _extract_text(path: Path) -> tuple[str, dict]:
    return path.read_text(encoding="utf-8", errors="replace"), {"extractor": "text"}


def _extract_pdf(path: Path) -> tuple[str, dict]:
    import fitz  # pymupdf
    doc = fitz.open(path)
    try:
        text = "\n\n".join(page.get_text() for page in doc)
        return text, {"extractor": "pdf", "pages": doc.page_count}
    finally:
        doc.close()


def _extract_docx(path: Path) -> tuple[str, dict]:
    from docx import Document
    doc = Document(path)
    text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return text, {"extractor": "docx"}


def _extract_pptx(path: Path) -> tuple[str, dict]:
    from pptx import Presentation
    prs = Presentation(path)
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text)
    return "\n\n".join(parts), {"extractor": "pptx", "slides": len(prs.slides)}


def _extract_mp4(path: Path) -> tuple[str, dict]:
    raise NotYetImplemented("MP4 transcription not yet implemented")


_DISPATCH = {
    ".txt": _extract_text, ".md": _extract_text, ".rst": _extract_text,
    ".py": _extract_text, ".sql": _extract_text, ".csv": _extract_text,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".mp4": _extract_mp4,
}


def extract_text(path: Path) -> tuple[str, dict]:
    """Return (text, metadata). Raises UnsupportedFileType or NotYetImplemented."""
    ext = path.suffix.lower()
    extractor = _DISPATCH.get(ext)
    if extractor is None:
        raise UnsupportedFileType(f"unsupported file type: {ext}")
    return extractor(path)
